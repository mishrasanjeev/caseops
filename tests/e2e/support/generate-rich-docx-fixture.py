"""Regenerate tests/fixtures/matter-rich-preview.docx with python-docx."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def main() -> None:
    document = Document()
    document.add_heading("CaseOps rich DOCX acceptance", level=1)
    paragraph = document.add_paragraph("Ordinary paragraph with ")
    paragraph.add_run("bold evidence").bold = True
    paragraph.add_run(", ")
    paragraph.add_run("italic analysis").italic = True
    paragraph.add_run(", and ")
    paragraph.add_run("underlined citation").underline = True
    paragraph.add_run(".")
    document.add_paragraph("First bullet item", style="List Bullet")
    document.add_paragraph("Second bullet item", style="List Bullet")
    document.add_paragraph("First numbered item", style="List Number")
    document.add_paragraph("Second numbered item", style="List Number")
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    for row, values in zip(table.rows, [("Issue", "Evidence"), ("DOCX", "Preserved")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    aligned = document.add_paragraph("Right aligned conclusion")
    aligned.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    aligned.paragraph_format.space_before = Pt(12)
    document.add_page_break()
    document.add_heading("Second page heading", level=2)
    destination = Path(__file__).resolve().parents[2] / "fixtures" / "matter-rich-preview.docx"
    document.save(destination)


if __name__ == "__main__":
    main()
