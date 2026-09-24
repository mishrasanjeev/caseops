# The template's exact user-facing error strings and field maps are intentionally
# kept readable; ruff's 100-column rule is not useful for those bounded literals.
# ruff: noqa: E501

from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import zipfile
from datetime import date, datetime
from itertools import islice
from typing import Any

from fastapi import HTTPException, status
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session, joinedload

from caseops_api.db.models import (
    CompanyMembership,
    Matter,
    MatterBulkUpdateOperation,
    Team,
)
from caseops_api.schemas.matter_bulk_updates import (
    MatterBulkUpdateApplyResponse,
    MatterBulkUpdateHistoryRecord,
    MatterBulkUpdateHistoryResponse,
    MatterBulkUpdatePreviewResponse,
    MatterBulkUpdateRow,
    MatterBulkUpdateSummary,
)
from caseops_api.schemas.matters import MatterUpdateRequest, normalize_matter_code
from caseops_api.services.audit import record_from_context
from caseops_api.services.matter_access import assert_access, can_access
from caseops_api.services.matters import update_matter
from caseops_api.services.session_context import SessionContext

HEADERS = [
    "Matter Title",
    "Matter Code",
    "Matter Type",
    "Practice Area",
    "Matter Description",
    "Client Name",
    "Client Code",
    "Client Contact Number",
    "Client Email",
    "Opposing Party Name",
    "Opposing Counsel",
    "Forum",
    "Court",
    "Court Forum Number",
    "Case Number",
    "Filing Number",
    "Filing Date",
    "Matter Owner",
    "Assigned Team",
    "Responsible Lawyer",
    "Temporary E-Case Number",
    "CNR Number",
    "Next Hearing Date",
]
MAX_ROWS = 500
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_XLSX_EXPANDED_BYTES = 16 * 1024 * 1024
MAX_XLSX_PART_BYTES = 8 * 1024 * 1024
MAX_XLSX_PARTS = 128
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_TYPE = "text/csv; charset=utf-8"


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _json_value(value: Any) -> object:
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def _parse_date(value: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise ValueError("Filing Date must be an ISO date such as 2026-09-21.") from exc


def _check_xlsx_archive(content: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_XLSX_PARTS or sum(item.file_size for item in entries) > MAX_XLSX_EXPANDED_BYTES:
            raise HTTPException(status_code=400, detail="The XLSX workbook is too large after expansion.")
        expanded = 0
        for item in entries:
            if item.file_size > MAX_XLSX_PART_BYTES:
                raise HTTPException(status_code=400, detail="The XLSX workbook contains an oversized part.")
            part_size = 0
            with archive.open(item) as stream:
                while chunk := stream.read(64 * 1024):
                    part_size += len(chunk)
                    expanded += len(chunk)
                    if part_size > MAX_XLSX_PART_BYTES or expanded > MAX_XLSX_EXPANDED_BYTES:
                        raise HTTPException(status_code=400, detail="The XLSX workbook is too large after expansion.")


def _parse_rows(content: bytes, filename: str) -> tuple[str, list[tuple[int, dict[str, str]]], str]:
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"Bulk update file must be {MAX_FILE_BYTES // (1024 * 1024)} MiB or smaller.",
        )
    extension = os.path.splitext(os.path.basename(filename))[1].casefold()
    if extension not in {".xlsx", ".csv"}:
        raise HTTPException(status_code=400, detail="Choose a .xlsx or .csv bulk-update file.")
    manifest_format = extension[1:]
    try:
        if manifest_format == "csv":
            text = content.decode("utf-8-sig", errors="strict")
            reader = csv.reader(io.StringIO(text, newline=""), strict=True)
            values = []
            for row in reader:
                values.append(row)
                if len(values) > MAX_ROWS + 1:
                    raise HTTPException(
                        status_code=400,
                        detail=f"Bulk update supports at most {MAX_ROWS} data rows.",
                    )
        else:
            _check_xlsx_archive(content)
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            try:
                sheet = workbook.active
                if sheet.max_column is not None and sheet.max_column > len(HEADERS):
                    raise HTTPException(
                        status_code=400,
                        detail="The workbook contains columns outside the bulk update template.",
                    )
                if sheet.max_row is not None and sheet.max_row > MAX_ROWS + 1:
                    raise HTTPException(
                        status_code=400,
                        detail=f"Bulk update supports at most {MAX_ROWS} data rows.",
                    )
                values = []
                for row_number, row in enumerate(
                    islice(sheet.iter_rows(max_col=len(HEADERS)), MAX_ROWS + 2),
                    start=1,
                ):
                    if any(cell.data_type == "f" for cell in row):
                        raise HTTPException(
                            status_code=400,
                            detail=f"Spreadsheet formulas are not accepted (row {row_number}). Use plain values.",
                        )
                    values.append([cell.value for cell in row])
                if len(values) > MAX_ROWS + 1:
                    raise HTTPException(
                        status_code=400,
                        detail=f"Bulk update supports at most {MAX_ROWS} data rows.",
                    )
            finally:
                workbook.close()
    except HTTPException:
        raise
    except Exception as exc:  # noqa: BLE001
        label = "CSV file" if manifest_format == "csv" else "XLSX workbook"
        raise HTTPException(status_code=400, detail=f"Upload a valid UTF-8 {label}.") from exc
    if not values:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail="The workbook is empty."
        )
    received_headers = [_cell(value) for value in values[0]]
    if received_headers != HEADERS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={
                "message": "The bulk update template headers must match exactly.",
                "expected": HEADERS,
                "received": received_headers,
            },
        )
    rows: list[tuple[int, dict[str, str]]] = []
    for row_number, raw in enumerate(values[1:], start=2):
        if len(raw) > len(HEADERS) and any(_cell(value) for value in raw[len(HEADERS) :]):
            raise HTTPException(
                status_code=400,
                detail=f"Row {row_number} contains columns outside the bulk update template.",
            )
        cells = [_cell(value) for value in raw[: len(HEADERS)]]
        if any(value.startswith(("=", "+", "-", "@")) for value in cells):
            raise HTTPException(
                status_code=400,
                detail=f"Formula-like cell content is not accepted (row {row_number}). Use plain values.",
            )
        if not any(cells):
            continue
        if len(rows) >= MAX_ROWS:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Bulk update supports at most {MAX_ROWS} populated rows.",
            )
        rows.append((row_number, dict(zip(HEADERS, cells, strict=False))))
    return hashlib.sha256(content).hexdigest(), rows, manifest_format


def _forum_value(value: str) -> str:
    normalized = " ".join(value.lower().replace("-", " ").replace("_", " ").split())
    mapping = {
        "lower court": "lower_court",
        "district court": "lower_court",
        "high court": "high_court",
        "supreme court": "supreme_court",
        "tribunal": "tribunal",
        "arbitration": "arbitration",
        "advisory": "advisory",
    }
    return mapping.get(normalized, "")


def _resolve_membership(session: Session, *, company_id: str, label: str) -> str | None:
    needle = label.casefold()
    memberships = session.scalars(
        select(CompanyMembership)
        .options(joinedload(CompanyMembership.user))
        .where(CompanyMembership.company_id == company_id, CompanyMembership.is_active.is_(True))
    )
    matches = [
        membership
        for membership in memberships
        if membership.user.is_active
        and needle in {membership.user.email.casefold(), membership.user.full_name.casefold()}
    ]
    if len(matches) == 1:
        return matches[0].id
    if not matches:
        raise ValueError(f"No active company member matches '{label}'.")
    raise ValueError(f"'{label}' matches multiple active company members.")


def _resolve_team(session: Session, *, company_id: str, label: str) -> str:
    teams = session.scalars(
        select(Team).where(Team.company_id == company_id, Team.is_active.is_(True))
    )
    matches = [team for team in teams if team.name.casefold() == label.casefold()]
    if len(matches) != 1:
        raise ValueError(f"Assigned Team '{label}' must match one active team exactly.")
    return matches[0].id


def _updates_for_row(
    session: Session, *, context: SessionContext, matter: Matter, row: dict[str, str]
) -> tuple[dict[str, Any], list[str]]:
    updates: dict[str, Any] = {}
    errors: list[str] = []
    mapping = {
        "Matter Title": "title",
        "Matter Type": "matter_type",
        "Practice Area": "practice_area",
        "Matter Description": "description",
        "Client Name": "client_name",
        "Client Code": "client_code",
        "Client Contact Number": "client_contact_number",
        "Client Email": "client_email",
        "Opposing Party Name": "opposing_party",
        "Opposing Counsel": "opposing_counsel",
        "Court": "court_name",
        "Court Forum Number": "court_forum_number",
        "Case Number": "case_number",
        "Filing Number": "filing_number",
        "Temporary E-Case Number": "temporary_e_case_number",
        "CNR Number": "cnr_number",
        "Next Hearing Date": "next_hearing_on",
    }
    for source, target in mapping.items():
        if row[source]:
            updates[target] = row[source]
    if row["Forum"]:
        forum_value = _forum_value(row["Forum"])
        if not forum_value:
            errors.append("Forum must match a supported forum level.")
        else:
            updates["forum_level"] = forum_value
    for source, target in (
        ("Filing Date", "filing_date"),
        ("Next Hearing Date", "next_hearing_on"),
    ):
        if row[source]:
            try:
                updates[target] = _parse_date(row[source])
            except ValueError as exc:
                errors.append(str(exc))
    for source, target in (
        ("Matter Owner", "assignee_membership_id"),
        ("Responsible Lawyer", "responsible_lawyer_membership_id"),
    ):
        if row[source]:
            try:
                updates[target] = _resolve_membership(
                    session, company_id=context.company.id, label=row[source]
                )
            except ValueError as exc:
                errors.append(str(exc))
    if row["Assigned Team"]:
        try:
            updates["team_id"] = _resolve_team(
                session, company_id=context.company.id, label=row["Assigned Team"]
            )
        except ValueError as exc:
            errors.append(str(exc))
    if errors:
        return {}, errors
    try:
        validated = MatterUpdateRequest(
            expected_updated_at=matter.updated_at,
            **updates,
        )
    except Exception as exc:  # noqa: BLE001
        errors.append(str(exc).split(" [type=")[0])
        return {}, errors
    validated_updates = validated.model_dump(exclude_unset=True)
    validated_updates.pop("expected_updated_at", None)
    return validated_updates, []


FORUM_RESULT_FIELDS = {
    "forum_level", "court_id", "court_name", "forum_catalog_entry_id",
    "forum_state", "forum_district", "forum_city", "forum_consumer_level",
}


def _planned_fields(updates: dict[str, Any]) -> set[str]:
    fields = set(updates)
    if {"forum_level", "court_name"} & fields:
        fields.update(FORUM_RESULT_FIELDS)
    return fields


def _planned_changes(
    session: Session, *, context: SessionContext, matter: Matter, updates: dict[str, Any]
) -> dict[str, dict[str, object]]:
    if not updates:
        return {}
    fields = _planned_fields(updates)
    before = {field: _json_value(getattr(matter, field)) for field in fields}
    savepoint = session.begin_nested()
    try:
        projected = update_matter(
            session,
            context=context,
            matter_id=matter.id,
            payload=MatterUpdateRequest(expected_updated_at=matter.updated_at, **updates),
            commit=False,
            commit_access_denial=False,
        )
        after = {field: _json_value(getattr(projected, field)) for field in fields}
    finally:
        if savepoint.is_active:
            savepoint.rollback()
    return {
        field: {"old": before[field], "new": after[field]}
        for field in fields
        if before[field] != after[field]
    }


def _preview_rows(
    session: Session, *, context: SessionContext, rows: list[tuple[int, dict[str, str]]]
) -> list[MatterBulkUpdateRow]:
    codes: list[str] = []
    for _, row in rows:
        raw_code = row.get("Matter Code", "")
        if not raw_code:
            continue
        try:
            codes.append(normalize_matter_code(raw_code))
        except Exception:
            continue
    matters = session.scalars(
        select(Matter).where(Matter.company_id == context.company.id, Matter.matter_code.in_(codes))
    )
    by_code = {matter.matter_code: matter for matter in matters}
    code_counts: dict[str, int] = {}
    for _, row in rows:
        raw_code = row.get("Matter Code", "")
        try:
            code = normalize_matter_code(raw_code)
        except Exception:
            continue
        code_counts[code] = code_counts.get(code, 0) + 1
    result: list[MatterBulkUpdateRow] = []
    for row_number, row in rows:
        raw_code = row.get("Matter Code", "")
        try:
            code = normalize_matter_code(raw_code)
        except Exception:
            code = ""
        matter = by_code.get(code)
        if code and code_counts.get(code, 0) > 1:
            result.append(
                MatterBulkUpdateRow(
                    row_number=row_number,
                    matter_code=code,
                    status="invalid",
                    errors=[
                        "Matter Code appears more than once in this file; remove duplicate rows."
                    ],
                )
            )
            continue
        if not raw_code or matter is None:
            result.append(
                MatterBulkUpdateRow(
                    row_number=row_number,
                    matter_code=code or raw_code,
                    status="invalid",
                    errors=[
                        "Matter Code must match one existing matter; this workflow never creates matters."
                    ],
                )
            )
            continue
        try:
            assert_access(session, context=context, matter=matter)
        except HTTPException as exc:
            if exc.status_code != status.HTTP_404_NOT_FOUND:
                raise
            result.append(
                MatterBulkUpdateRow(
                    row_number=row_number,
                    matter_code=code,
                    status="invalid",
                    errors=[
                        "Matter Code must match one existing matter; this workflow never creates matters."
                    ],
                )
            )
            continue
        updates, errors = _updates_for_row(session, context=context, matter=matter, row=row)
        changes: dict[str, dict[str, object]] = {}
        if not errors:
            try:
                changes = _planned_changes(session, context=context, matter=matter, updates=updates)
            except HTTPException as exc:
                detail = exc.detail
                errors.append(
                    detail
                    if isinstance(detail, str)
                    else str(detail.get("message", "Invalid court selection."))
                )
        result.append(
            MatterBulkUpdateRow(
                row_number=row_number,
                matter_id=matter.id,
                matter_code=matter.matter_code,
                status="invalid" if errors else ("changed" if changes else "unchanged"),
                errors=errors,
                changes=changes,
                expected_updated_at=matter.updated_at,
            )
        )
    return result


def preview_matter_bulk_update(
    session: Session, *, context: SessionContext, content: bytes, filename: str
) -> MatterBulkUpdatePreviewResponse:
    file_hash, rows, _ = _parse_rows(content, filename)
    plans = _preview_rows(session, context=context, rows=rows)
    token = _preview_token(file_hash, plans)
    changed = sum(plan.status == "changed" for plan in plans)
    unchanged = sum(plan.status == "unchanged" for plan in plans)
    invalid = sum(plan.status == "invalid" for plan in plans)
    return MatterBulkUpdatePreviewResponse(
        preview_token=token,
        headers=HEADERS,
        summary=MatterBulkUpdateSummary(
            total_rows=len(plans),
            matched_rows=len(plans) - invalid,
            changed_rows=changed,
            unchanged_rows=unchanged,
            invalid_rows=invalid,
        ),
        rows=plans,
    )


def apply_matter_bulk_update(
    session: Session,
    *,
    context: SessionContext,
    content: bytes,
    filename: str,
    preview_token: str,
) -> MatterBulkUpdateApplyResponse:
    file_hash, rows, manifest_format = _parse_rows(content, filename)
    plans = _preview_rows(session, context=context, rows=rows)
    current_token = _preview_token(file_hash, plans)
    safe_filename = (
        os.path.basename(filename).replace("\\", "_").replace("/", "_")[:255] or "upload"
    )
    if current_token != preview_token:
        changed = sum(plan.status == "changed" for plan in plans)
        invalid = sum(plan.status == "invalid" for plan in plans)
        session.add(
            MatterBulkUpdateOperation(
                company_id=context.company.id,
                uploader_membership_id=context.membership.id,
                filename=safe_filename,
                format=manifest_format,
                status="stale",
                total_rows=len(plans),
                changed_rows=changed,
                invalid_rows=invalid,
                applied_rows=0,
                failed_rows=changed,
            )
        )
        session.commit()
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="A Matter changed after preview or the upload differs; review a fresh preview before applying.",
        )
    applied = 0
    work = sorted(zip(plans, rows, strict=True), key=lambda item: item[0].matter_id or "")
    applying_matter_id: str | None = None
    try:
        for plan, (_, raw_row) in work:
            if plan.status != "changed":
                continue
            applying_matter_id = plan.matter_id
            matter = session.get(Matter, plan.matter_id)
            if matter is None:
                raise HTTPException(status_code=409, detail="Matter no longer exists.")
            updates, errors = _updates_for_row(session, context=context, matter=matter, row=raw_row)
            if errors:
                raise ValueError("; ".join(errors))
            fields = _planned_fields(updates)
            before = {field: _json_value(getattr(matter, field)) for field in fields}
            updates["expected_updated_at"] = plan.expected_updated_at
            updated = update_matter(
                session,
                context=context,
                matter_id=matter.id,
                payload=MatterUpdateRequest(**updates),
                commit=False,
                commit_access_denial=False,
            )
            if any(
                _json_value(getattr(updated, field))
                != plan.changes.get(field, {}).get("new", before[field])
                for field in fields
            ):
                raise HTTPException(status_code=409, detail="Matter source state changed after preview.")
            plan.status = "applied"
            applied += 1
        skipped = sum(plan.status in {"invalid", "unchanged"} for plan in plans)
        operation = MatterBulkUpdateOperation(
            company_id=context.company.id,
            uploader_membership_id=context.membership.id,
            filename=safe_filename,
            format=manifest_format,
            status="completed_with_errors"
            if any(plan.status == "invalid" for plan in plans)
            else "completed",
            total_rows=len(plans),
            changed_rows=sum(bool(plan.changes) for plan in plans),
            invalid_rows=sum(plan.status == "invalid" for plan in plans),
            applied_rows=applied,
            failed_rows=0,
        )
        session.add(operation)
        session.commit()
    except Exception as exc:  # noqa: BLE001
        session.rollback()
        if not isinstance(exc, (HTTPException, ValueError)):
            raise
        if applying_matter_id and isinstance(exc, HTTPException) and exc.status_code == 404:
            denied_matter = session.get(Matter, applying_matter_id)
            if (
                denied_matter is not None
                and denied_matter.company_id == context.company.id
                and not can_access(session, context=context, matter=denied_matter)
            ):
                record_from_context(
                    session,
                    context,
                    action="access_denied",
                    target_type="matter",
                    target_id=denied_matter.id,
                    matter_id=denied_matter.id,
                    result="denied",
                    metadata={"reason": "matter_visibility_denied"},
                )
        session.add(
            MatterBulkUpdateOperation(
                company_id=context.company.id,
                uploader_membership_id=context.membership.id,
                filename=safe_filename,
                format=manifest_format,
                status="stale",
                total_rows=len(plans),
                changed_rows=sum(plan.status in {"changed", "applied"} for plan in plans),
                invalid_rows=sum(plan.status == "invalid" for plan in plans),
                applied_rows=0,
                failed_rows=sum(plan.status in {"changed", "applied"} for plan in plans),
            )
        )
        session.commit()
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="A Matter or its access changed while applying. No rows were updated; review a fresh preview.",
        ) from exc
    return MatterBulkUpdateApplyResponse(
        preview_token=preview_token,
        total_rows=len(plans),
        applied_rows=applied,
        skipped_rows=skipped,
        failed_rows=0,
        rows=plans,
        operation_id=operation.id,
    )


def _preview_token(file_hash: str, plans: list[MatterBulkUpdateRow]) -> str:
    state = [
        {
            "row": plan.row_number,
            "matter_id": plan.matter_id,
            "status": plan.status,
            "expected_updated_at": plan.expected_updated_at.isoformat()
            if plan.expected_updated_at
            else None,
            "changes": plan.changes,
            "errors": plan.errors,
        }
        for plan in plans
    ]
    payload = json.dumps([file_hash, state], sort_keys=True, separators=(",", ":"), default=str)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def matter_bulk_update_template(manifest_format: str) -> tuple[bytes, str, str]:
    if manifest_format == "csv":
        output = io.StringIO(newline="")
        csv.writer(output, lineterminator="\r\n").writerow(HEADERS)
        return output.getvalue().encode("utf-8"), CSV_TYPE, "matter-bulk-update-template.csv"
    if manifest_format != "xlsx":
        raise HTTPException(status_code=400, detail="Format must be csv or xlsx.")
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Matter Updates"
    sheet.append(HEADERS)
    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Bulk update existing matters"])
    instructions.append(
        ["Matter Code is required and must match one existing matter in this workspace."]
    )
    instructions.append(["Blank cells leave the existing value unchanged."])
    instructions.append(
        ["Matter Status is deliberately excluded; use the dedicated lifecycle workflow."]
    )
    instructions.append(
        ["Use plain values only. Formulas are rejected. Maximum 500 rows and 5 MiB."]
    )
    instructions.append(["Review every old/new value and invalid row before applying changes."])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue(), XLSX_TYPE, "matter-bulk-update-template.xlsx"


def list_matter_bulk_update_history(
    session: Session, *, context: SessionContext, limit: int = 50
) -> MatterBulkUpdateHistoryResponse:
    query = select(MatterBulkUpdateOperation).where(
        MatterBulkUpdateOperation.company_id == context.company.id
    )
    total = session.scalar(select(func.count()).select_from(query.subquery())) or 0
    operations = session.scalars(
        query.options(
            joinedload(MatterBulkUpdateOperation.uploader_membership).joinedload(
                CompanyMembership.user
            )
        )
        .order_by(MatterBulkUpdateOperation.created_at.desc())
        .limit(limit)
    )
    records = []
    for item in operations:
        membership = item.uploader_membership
        records.append(
            MatterBulkUpdateHistoryRecord(
                id=item.id,
                filename=item.filename,
                format=item.format,
                status=item.status,
                total_rows=item.total_rows,
                changed_rows=item.changed_rows,
                invalid_rows=item.invalid_rows,
                applied_rows=item.applied_rows,
                failed_rows=item.failed_rows,
                uploader_membership_id=item.uploader_membership_id,
                uploader_name=membership.user.full_name if membership else None,
                uploader_email=membership.user.email if membership else None,
                created_at=item.created_at,
            )
        )
    return MatterBulkUpdateHistoryResponse(operations=records, total=total)
