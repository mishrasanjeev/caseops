"""IPLF-032 / UJ-02: source-tolerant portfolio imports without legal fabrication."""

from __future__ import annotations

import csv
import io
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, date, datetime, timedelta
from hashlib import sha256
from threading import Event
from time import monotonic, sleep
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import event, select, text

from caseops_api.db.models import (
    Company,
    CompanyMembership,
    InAppNotification,
    IpDocketRecord,
    IpRenewalTerm,
    IpTrademarkParticularVersion,
    NotificationDeliveryIntent,
)
from caseops_api.db.session import get_engine, get_session_factory
from caseops_api.services import ip_imports, ip_renewals, notification_delivery
from caseops_api.services.ip_import_files import parse_ip_import_file
from tests.sql_observation import is_read_only_select
from tests.test_auth_company import auth_headers, bootstrap_company


def test_compressed_xlsx_cannot_amplify_staged_source_text_unboundedly():
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet()
    sheet.append(["Source one", "Source two", "Source three"])
    for _ in range(500):
        sheet.append(["x" * 10000] * 3)
    content = io.BytesIO()
    workbook.save(content)
    workbook.close()
    assert len(content.getvalue()) < 10 * 1024 * 1024
    with pytest.raises(HTTPException, match="Decoded source text exceeds 10 MB"):
        parse_ip_import_file(filename="compressed-source.xlsx", content=content.getvalue())


def test_dimensionless_xlsx_remains_bounded_and_retains_native_date():
    workbook = Workbook()
    workbook.active.append(["Trademark", "Application Date"])
    workbook.active.append(["Dimensionless source", datetime(2026, 1, 31)])
    original, dimensionless = io.BytesIO(), io.BytesIO()
    workbook.save(original)
    workbook.close()
    with ZipFile(original) as reader, ZipFile(dimensionless, "w") as writer:
        for name in reader.namelist():
            content = reader.read(name)
            if name == "xl/worksheets/sheet1.xml":
                sheet = ElementTree.fromstring(content)
                dimension = sheet.find(
                    "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}dimension"
                )
                sheet.remove(dimension)
                content = ElementTree.tostring(sheet)
            writer.writestr(name, content)
    parsed = parse_ip_import_file(filename="dimensionless.xlsx", content=dimensionless.getvalue())
    assert len(parsed.rows) == 1
    assert parsed.rows[0].values["Application Date"] == "2026-01-31T00:00:00"


def test_native_xlsx_dates_stage_as_json_safe_source_evidence(client: TestClient) -> None:
    actor = bootstrap_company(client)
    headers = auth_headers(str(actor["access_token"]))
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Title", "Mark Text", "Nice Class", "Applicant", "Application Date", "Duration"])
    sheet.append(
        [
            "SOURCE DATE",
            "SOURCE DATE",
            9,
            "Source Owner LLP",
            datetime(2026, 1, 31),
            timedelta(hours=26),
        ]
    )
    sheet["F2"].number_format = "[h]:mm:ss"
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    with TestClient(client.app, raise_server_exceptions=False) as transport:
        response = transport.post(
            "/api/ip/imports/upload",
            headers={
                **headers,
                "Origin": "http://localhost:3000",
                "X-Request-ID": "ip-native-date-baseline-20261013",
                "X-CaseOps-Automated-Test": "no-paid-providers",
            },
            files={
                "file": (
                    "native-dates.xlsx",
                    stream.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert response.status_code == 201, response.text
    assert response.headers["access-control-allow-origin"] == "http://localhost:3000"
    assert response.json()["rows"][0]["normalized"]["application_date"] == "2026-01-31"
    assert (
        response.json()["rows"][0]["normalized"]["_source_values"]["Duration"] == "1 day, 2:00:00"
    )


def _actor(client: TestClient, *, configure: bool = True) -> dict:
    actor = bootstrap_company(client)
    actor["headers"] = {
        **auth_headers(str(actor["access_token"])),
        "X-CaseOps-Automated-Test": "no-paid-providers",
    }
    if configure:
        response = client.put(
            "/api/ip/workspace/configuration",
            headers=actor["headers"],
            json={
                "enabled_asset_types": ["trademark"],
                "jurisdictions": ["IN"],
                "offices": ["IP India"],
                "timezone": "Asia/Kolkata",
                "holiday_calendar_key": "source-fixture",
                "working_day_policy": {
                    "working_weekdays": [0, 1, 2, 3, 4],
                },
                "document_taxonomy_version": "fixture-v1",
                "event_catalog_version": "fixture-v1",
                "deadline_rule_versions": {},
                "notification_channels": ["in_app"],
                "critical_event_policy": {"escalation_after_minutes": 60},
                "escalation_owner_membership_id": actor["membership"]["id"],
            },
        )
        assert response.status_code == 200, response.text
    return actor


def _upload(client, actor, headers, rows):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(headers)
    writer.writerows(rows)
    content = stream.getvalue().encode()
    response = client.post(
        "/api/ip/imports/upload",
        headers=actor["headers"],
        files={"file": ("source.csv", content, "text/csv")},
    )
    assert response.status_code == 201, response.text
    preview = response.json()
    assert preview["job"]["source_sha256"] == sha256(content).hexdigest()
    return preview


def _commit(client, actor, preview, key="source-import-key-20261013"):
    started = monotonic()
    response = client.post(
        f"/api/ip/imports/{preview['job']['id']}/commit",
        headers=actor["headers"],
        json={
            "preview_token": preview["job"]["preview_token"],
            "idempotency_key": key,
        },
    )
    result = response.json()
    assert monotonic() - started < 90, "A bounded commit must fit the existing browser deadline."
    limit = result["job"]["commit_batch_size"] if response.status_code == 202 else 50
    assert response.status_code in {200, 202}, response.text
    for _ in range((preview["job"]["total_rows"] + limit - 1) // limit):
        if response.status_code == 200:
            assert result["job"]["status"] in {"committed", "committed_with_errors"}
            return result
        assert result["job"]["status"] == "staged"
        pending = sum(row["commit_status"] == "pending" for row in result["rows"])
        version = result["job"]["version"]
        assert client.get("/api/health", headers=actor["headers"]).status_code == 200
        started = monotonic()
        response = client.post(
            f"/api/ip/imports/{preview['job']['id']}/resume",
            headers=actor["headers"],
            json={"expected_job_version": version, "preview_token": result["job"]["preview_token"]},
        )
        assert response.status_code in {200, 202}, response.text
        assert monotonic() - started < 90, (
            "A bounded resume must fit the existing browser deadline."
        )
        result = response.json()
        assert result["job"]["version"] > version
        assert sum(row["commit_status"] == "pending" for row in result["rows"]) < pending
    assert response.status_code == 200, response.text
    return result


def test_derived_display_title_is_admitted_before_canonical_materialization(client):
    actor = _actor(client)
    preview = _upload(client, actor, ["Application No."], [["1"], ["12"]])
    assert preview["rows"][0]["validation_status"] == "invalid"
    assert {"field": "title", "code": "invalid_length"} in preview["rows"][0]["errors"]
    saved = _commit(client, actor, preview)
    assert saved["job"]["committed_rows"] == 1
    assert saved["job"]["failed_rows"] == 0
    assert [row["commit_status"] for row in saved["rows"]] == ["skipped", "committed"]


def test_supplied_headers_partial_commit_and_no_fabricated_legal_state(client, monkeypatch):
    actor = _actor(client)
    monkeypatch.setattr(ip_renewals, "_now", lambda: datetime(2028, 1, 1, tzinfo=UTC))
    headers = [
        "S.No.",
        "Trademark",
        "Trademark Image",
        "Application Date",
        "Renewal Due Date",
        "Applicant Name(s)",
        "Class",
        "Application No.",
        "Status",
        "Comment",
        "Takeover request",
        "Name Change request",
        "Address Change request",
        "Internal Remarks",
        "Unknown extra",
    ]
    positive = [
        "source-7",
        "SOURCE MARK",
        "logo-reference",
        "31/01/2026",
        "31/03/2028",
        "Owner One; Owner Two",
        "9, 36, and 42",
        "TM/2028/7",
        "Registered",
        "comment",
        "Yes",
        "Pending",
        "N/A",
        "privileged source note",
        "retained extra",
    ]
    warning = ["source-8", "WARNING MARK", "", "bad date", "N.A.", "", "", "TM/2028/8"]
    unusable = ["source-9", "", "", "", "", "", "", ""]
    preview = _upload(client, actor, headers, [positive, positive, warning, unusable, [""] * 15])
    assert (
        preview["job"]["total_rows"],
        preview["job"]["valid_rows"],
        preview["job"]["invalid_rows"],
    ) == (5, 3, 2)
    first, duplicate, warned, bad, blank = preview["rows"]
    assert first["normalized"]["class_numbers"] == [9, 36, 42]
    assert first["normalized"]["applicant_name"] == "Owner One; Owner Two"
    assert first["normalized"]["source_status"] == "Registered"
    assert first["normalized"]["source_reference"] == "source-7"
    assert first["normalized"]["_source_values"]["Unknown extra"] == "retained extra"
    assert first["normalized"]["renewal_due_date"] == "2028-03-31"
    assert "address_change_request" not in first["normalized"]
    assert duplicate["reconciliation_decision"] == "skip"
    assert "renewal_due_date" not in warned["normalized"]
    assert "application_date" not in warned["normalized"]
    assert {"field": "application_date", "code": "invalid_optional_date"} in (
        warned["normalized"]["_warnings"]
    )
    assert {e["code"] for e in bad["errors"]} == {"unusable_row"}
    assert {e["code"] for e in blank["errors"]} == {"blank_row"}
    with get_session_factory()() as session:
        assert session.scalar(select(IpDocketRecord.id)) is None
    committed = _commit(client, actor, preview)
    assert committed["job"]["status"] == "committed"
    assert committed["job"]["committed_rows"] == 2
    assert [r["commit_status"] for r in committed["rows"]] == [
        "committed",
        "skipped",
        "committed",
        "skipped",
        "skipped",
    ]
    docket_id = committed["rows"][0]["created_docket_id"]
    docket = client.get(f"/api/ip/dockets/{docket_id}", headers=actor["headers"]).json()
    assert docket["status"] == "draft"
    assert docket["current_particulars"]["readiness_status"] == "incomplete"
    assert [c["specification"] for c in docket["current_particulars"]["classes_json"]] == [
        "",
        "",
        "",
    ]
    assert "privileged source note" not in str(docket)
    portfolio = client.get("/api/ip/portfolio", headers=actor["headers"]).json()
    assert len(portfolio["rows"]) == 2
    assert "privileged source note" not in str(portfolio)
    with get_session_factory()() as session:
        assert session.scalar(select(IpRenewalTerm.id)) is None
        intents = session.scalars(select(NotificationDeliveryIntent)).all()
        assert len(intents) == 1
        assert ip_renewals._aware(intents[0].scheduled_for) == datetime(
            2028, 2, 29, 3, 30, tzinfo=UTC
        )
        assert intents[0].recipient_membership_id == actor["membership"]["id"]
        assert "not a verified legal deadline" in intents[0].body
    repeated = _commit(client, actor, preview)
    assert repeated["replayed"] is True
    assert repeated["rows"] == committed["rows"]
    persisted = client.get(
        f"/api/ip/imports/{preview['job']['id']}", headers=actor["headers"]
    ).json()
    assert persisted["rows"] == committed["rows"]


def test_ambiguous_source_identity_and_invalid_rows_do_not_reserve_duplicate_keys(client):
    actor = _actor(client)
    preview = _upload(
        client,
        actor,
        ["  TRADEMARK ", "App. No.", "Application Number", "Class"],
        [
            ["Invalid first", "APP/1", "APP/2", "99"],
            ["Valid second", "APP/1", "APP/1", "9"],
            ["Valid third", "APP/1", "APP/1", "9"],
        ],
    )
    assert {e["code"] for e in preview["rows"][0]["errors"]} == {
        "ambiguous_headers",
        "out_of_range",
    }
    assert not preview["rows"][1]["duplicate_candidates"]
    assert preview["rows"][2]["reconciliation_decision"] == "skip"
    result = _commit(client, actor, preview)
    assert result["job"]["committed_rows"] == 1
    assert (
        client.get("/api/ip/portfolio", headers=actor["headers"]).json()["rows"][0]["asset_title"]
        == "Valid second"
    )


def test_missing_context_is_not_guessed_and_missing_mark_uses_only_real_number(client):
    actor = _actor(client, configure=False)
    blocked = _upload(client, actor, ["Trademark"], [["Context absent"]])
    assert {e["field"] for e in blocked["rows"][0]["errors"]} == {"jurisdiction", "office"}
    assert "jurisdiction" not in blocked["rows"][0]["normalized"]
    preview = _upload(
        client,
        actor,
        ["Application No.", "Jurisdiction", "Office"],
        [
            ["ACTUAL/1", "IN", "IP India"],
        ],
    )
    result = _commit(client, actor, preview)
    docket = client.get(
        f"/api/ip/dockets/{result['rows'][0]['created_docket_id']}", headers=actor["headers"]
    ).json()
    assert docket["title"] == "ACTUAL/1"
    assert "text" not in docket["current_particulars"]["representation_json"]
    assert docket["current_particulars"]["parties_json"] == []


def test_distinct_applications_with_same_mark_survive_and_late_duplicate_is_skipped(client):
    actor = _actor(client)
    headers = ["Trademark", "Application No.", "Applicant", "Class"]
    first = _upload(client, actor, headers, [["Same mark", "APP/31", "Owner", "9"]])
    stale = _upload(
        client,
        actor,
        headers,
        [["Same mark", "APP 31", "Owner", "9"], ["Same mark", "APP/32", "Owner", "9"]],
    )
    _commit(client, actor, first, "first-source-key")
    result = _commit(client, actor, stale, "second-source-key")
    assert result["rows"][0]["commit_error_code"] == "duplicate_skipped"
    assert result["rows"][1]["created_docket_id"]
    portfolio = client.get("/api/ip/portfolio", headers=actor["headers"]).json()
    assert {tuple(row["application_numbers"]) for row in portfolio["rows"]} == {
        ("APP/31",),
        ("APP/32",),
    }


def test_cautious_fallback_does_not_merge_same_mark_different_applicant(client):
    actor = _actor(client)
    preview = _upload(
        client,
        actor,
        ["Trademark", "Applicant", "Class"],
        [
            ["Fallback mark", "Owner One", "9,36"],
            ["Fallback mark", "Owner One", "36,9"],
            ["Fallback mark", "Owner Two", "9,36"],
        ],
    )
    result = _commit(client, actor, preview)
    assert result["job"]["committed_rows"] == 2
    assert result["rows"][1]["commit_status"] == "skipped"
    assert result["rows"][2]["created_docket_id"]


@pytest.mark.parametrize(
    ("due", "expected"),
    [
        (date(2028, 3, 31), datetime(2028, 2, 29, 3, 30, tzinfo=UTC)),
        (date(2027, 3, 31), datetime(2027, 2, 28, 3, 30, tzinfo=UTC)),
        (date(2028, 1, 31), datetime(2027, 12, 31, 3, 30, tzinfo=UTC)),
        (date(2028, 2, 29), datetime(2028, 1, 29, 3, 30, tzinfo=UTC)),
    ],
)
def test_calendar_month_alert_clamps_month_end(due, expected):
    assert ip_renewals.portfolio_renewal_alert_at(due, timezone="Asia/Kolkata") == expected


def test_source_alert_uses_tenant_timezone_without_workspace_configuration(client, monkeypatch):
    actor = _actor(client, configure=False)
    with get_session_factory()() as session:
        session.get(Company, actor["company"]["id"]).timezone = "America/New_York"
        session.commit()
    monkeypatch.setattr(ip_renewals, "_now", lambda: datetime(2035, 1, 1, tzinfo=UTC))
    preview = _upload(
        client,
        actor,
        ["Trademark", "Renewal Date", "Jurisdiction", "Office"],
        [["Tenant timezone source", "2035-04-30", "IN", "IP India"]],
    )
    row = _commit(client, actor, preview)["rows"][0]
    assert datetime.fromisoformat(
        row["normalized"]["_renewal_reminder"]["scheduled_for"]
    ) == datetime(
        2035,
        3,
        30,
        13,
        tzinfo=UTC,
    )


def test_due_soon_alert_uses_existing_delivery_owner_and_idempotency(client, monkeypatch):
    actor = _actor(client)
    now = datetime(2028, 3, 20, 0, tzinfo=UTC)
    monkeypatch.setattr(ip_renewals, "_now", lambda: now)
    preview = _upload(client, actor, ["Trademark", "Renewal Date"], [["Due soon", "2028-03-31"]])
    result = _commit(client, actor, preview)
    reminder = result["rows"][0]["normalized"]["_renewal_reminder"]
    assert datetime.fromisoformat(reminder["scheduled_for"]) == now
    assert reminder["status"] == "queued"
    assert _commit(client, actor, preview)["rows"][0]["normalized"]["_renewal_reminder"] == reminder


@pytest.mark.parametrize("active", [True, False], ids=["authorized-recipient", "revoked-recipient"])
def test_source_reminder_actual_in_app_dispatch_rechecks_recipient(client, monkeypatch, active):
    actor = _actor(client)
    now = datetime(2028, 3, 20, tzinfo=UTC)
    monkeypatch.setattr(ip_renewals, "_now", lambda: now)
    monkeypatch.setattr(notification_delivery, "_now", lambda: now)
    preview = _upload(
        client, actor, ["Trademark", "Renewal Date"], [["Dispatch source", "2028-03-31"]]
    )
    result = _commit(client, actor, preview)
    intent_id = result["rows"][0]["normalized"]["_renewal_reminder"]["intent_id"]
    with get_session_factory()() as session:
        intent = session.get(NotificationDeliveryIntent, intent_id)
        version = session.scalar(
            select(IpTrademarkParticularVersion).where(
                IpTrademarkParticularVersion.docket_id == result["rows"][0]["created_docket_id"],
                IpTrademarkParticularVersion.version == 1,
            )
        )
        assert intent.source_type == "ip_trademark_particular_version"
        assert intent.source_id == version.id
        assert ip_renewals.imported_renewal_reminder_is_current(session, intent=intent)
        membership = session.get(CompanyMembership, actor["membership"]["id"])
        if not active:
            membership.is_active = False
            session.commit()
        delivered = notification_delivery.process_notification_delivery_intent(
            session,
            intent_id=intent_id,
            company_id=actor["company"]["id"],
        )
        session.commit()
        session.expire_all()
        intent = session.get(NotificationDeliveryIntent, intent_id)
        assert str(delivered.status) == ("delivered" if active else "blocked")
        assert str(intent.status) == ("delivered" if active else "blocked")
        notifications = session.scalars(
            select(InAppNotification).where(
                InAppNotification.company_id == actor["company"]["id"],
                InAppNotification.source_type == "ip_trademark_particular_version",
            )
        ).all()
        assert len(notifications) == (1 if active else 0)
        if active:
            assert notifications[0].recipient_membership_id == actor["membership"]["id"]
            assert "2028-03-31" in notifications[0].body


@pytest.mark.parametrize("new_due", ["2028-04-30", None], ids=["date-change", "date-removal"])
def test_source_predicate_rejects_actual_canonical_version_correction(client, monkeypatch, new_due):
    """The real canonical writer cancels/remakes intents, not only the predicate."""
    from tests.test_ip_current_source_reminders_20261013 import _append_current

    actor = _actor(client)
    monkeypatch.setattr(ip_renewals, "_now", lambda: datetime(2028, 1, 1, tzinfo=UTC))
    dispatch_at = datetime(2028, 4, 20, tzinfo=UTC)
    monkeypatch.setattr(notification_delivery, "_now", lambda: dispatch_at)
    preview = _upload(
        client,
        actor,
        ["Trademark", "Renewal Date", "Applicant", "Class", "Goods/Services"],
        [["Correction source", "2028-03-31", "Source Owner", "9", "Legal software"]],
    )
    result = _commit(client, actor, preview)
    row = result["rows"][0]
    docket_id = row["created_docket_id"]
    corrected = _append_current(client, actor, docket_id, due=new_due)
    assert corrected["current_version"] == 3
    with get_session_factory()() as session:
        intent = session.get(
            NotificationDeliveryIntent, row["normalized"]["_renewal_reminder"]["intent_id"]
        )
        assert not ip_renewals.imported_renewal_reminder_is_current(session, intent=intent)
        assert str(intent.status) == "cancelled"
        # Retained legacy/repair state must not bypass the real source fence,
        # even when cancellation itself is deliberately taken out of the oracle.
        intent.status = "queued"
        intent.scheduled_for = dispatch_at
        intent.next_attempt_at = None
        session.commit()
        denied = notification_delivery.process_notification_delivery_intent(
            session,
            intent_id=intent.id,
            company_id=actor["company"]["id"],
        )
        session.commit()
        session.refresh(intent)
        assert str(denied.status) == "blocked"
        assert intent.dead_letter_reason == "ip_portfolio_renewal_source_stale"
        assert intent.in_app_notification_id is None and intent.attempts == 0
        assert (
            session.scalar(
                select(InAppNotification.id).where(
                    InAppNotification.company_id == actor["company"]["id"],
                    InAppNotification.event_type == "renewal_source_date_1month",
                )
            )
            is None
        )
        current = session.scalars(
            select(NotificationDeliveryIntent).where(
                NotificationDeliveryIntent.ip_docket_id == docket_id,
                NotificationDeliveryIntent.status == "queued",
            )
        ).all()
        assert len(current) == (1 if new_due else 0)
        if new_due:
            assert current[0].metadata_json["source_due_on"] == new_due
            assert ip_renewals.imported_renewal_reminder_is_current(session, intent=current[0])
            delivered = notification_delivery.process_notification_delivery_intent(
                session,
                intent_id=current[0].id,
                company_id=actor["company"]["id"],
            )
            session.commit()
            assert str(delivered.status) == "delivered"
            notes = session.scalars(
                select(InAppNotification).where(
                    InAppNotification.company_id == actor["company"]["id"],
                    InAppNotification.event_type == "renewal_source_date_1month",
                )
            ).all()
            assert len(notes) == 1
            assert new_due in notes[0].body and "2028-03-31" not in notes[0].body
    retained = client.get(
        f"/api/ip/imports/{preview['job']['id']}", headers=actor["headers"]
    ).json()
    assert retained["rows"][0]["normalized"]["renewal_due_date"] == "2028-03-31"


def test_source_date_edit_clear_stale_replay_and_actual_delivery(client, monkeypatch):
    actor = _actor(client)
    now = datetime(2028, 3, 20, tzinfo=UTC)
    monkeypatch.setattr(ip_renewals, "_now", lambda: now)
    monkeypatch.setattr(notification_delivery, "_now", lambda: now)
    preview = _upload(client, actor, ["Trademark", "Renewal Date"], [["Edit source", "2028-03-31"]])
    row = _commit(client, actor, preview)["rows"][0]
    docket_id = row["created_docket_id"]
    url = f"/api/ip/dockets/{docket_id}/portfolio-renewal-date"
    changed = client.patch(
        url,
        headers=actor["headers"],
        json={
            "expected_current_version": 1,
            "renewal_due_date": "2028-03-30",
            "reason": "Corrected against source date evidence",
        },
    )
    assert changed.status_code == 200, changed.text
    record = changed.json()
    assert record["current_version"] == 2
    assert record["status"] == "draft"
    assert record["current_particulars"]["parties_json"] == []
    assert record["current_particulars"]["classes_json"] == []
    stale = client.patch(
        url,
        headers=actor["headers"],
        json={
            "expected_current_version": 1,
            "renewal_due_date": None,
            "reason": "Stale correction",
        },
    )
    assert stale.status_code == 409
    with get_session_factory()() as session:
        old = session.get(
            NotificationDeliveryIntent, row["normalized"]["_renewal_reminder"]["intent_id"]
        )
        assert str(old.status) == "cancelled"
        old_replay = notification_delivery.process_notification_delivery_intent(
            session,
            intent_id=old.id,
            company_id=actor["company"]["id"],
        )
        assert str(old_replay.status) == "cancelled"
        pending = session.scalar(
            select(NotificationDeliveryIntent).where(
                NotificationDeliveryIntent.ip_docket_id == docket_id,
                NotificationDeliveryIntent.status == "queued",
            )
        )
        assert pending.source_id == record["current_particulars"]["id"]
        assert pending.metadata_json["source_due_on"] == "2028-03-30"
        delivery = notification_delivery.process_notification_delivery_intent(
            session,
            intent_id=pending.id,
            company_id=actor["company"]["id"],
        )
        assert str(delivery.status) == "delivered"
        session.commit()
        notes = session.scalars(
            select(InAppNotification).where(
                InAppNotification.company_id == actor["company"]["id"],
                InAppNotification.source_type == "ip_trademark_particular_version",
                InAppNotification.source_id == record["current_particulars"]["id"],
            )
        ).all()
        assert len(notes) == 1
        assert "2028-03-30" in notes[0].body
        assert "2028-03-31" not in notes[0].body
    # A second new date creates a pending version-specific reminder, then clear cancels it.
    changed = client.patch(
        url,
        headers=actor["headers"],
        json={
            "expected_current_version": 2,
            "renewal_due_date": "2028-04-30",
            "reason": "Second source date correction",
        },
    )
    assert changed.status_code == 200, changed.text
    cleared = client.patch(
        url,
        headers=actor["headers"],
        json={
            "expected_current_version": 3,
            "renewal_due_date": None,
            "reason": "Source date withdrawn",
        },
    )
    assert cleared.status_code == 200, cleared.text
    assert cleared.json()["current_version"] == 4
    assert (
        cleared.json()["current_particulars"]["representation_json"]["portfolio_import"][
            "renewal_due_date"
        ]
        is None
    )
    with get_session_factory()() as session:
        assert (
            session.scalar(
                select(NotificationDeliveryIntent.id).where(
                    NotificationDeliveryIntent.ip_docket_id == docket_id,
                    NotificationDeliveryIntent.status == "queued",
                )
            )
            is None
        )
        assert (
            session.scalar(select(IpRenewalTerm.id).where(IpRenewalTerm.docket_id == docket_id))
            is None
        )
    retained = client.get(
        f"/api/ip/imports/{preview['job']['id']}", headers=actor["headers"]
    ).json()
    assert retained["rows"][0]["normalized"]["renewal_due_date"] == "2028-03-31"

    closed = client.post(
        f"/api/ip/dockets/{docket_id}/lifecycle",
        headers=actor["headers"],
        json={
            "expected_lifecycle_version": 0,
            "to_status": "closed",
            "effective_at": datetime.now(UTC).isoformat(),
            "reason": "Synthetic source record closed",
            "outcome": "closed",
            "source": "source_review",
            "evidence_ref": "synthetic:source-close",
            "linked_matter_handling": "not_linked",
        },
    )
    assert closed.status_code == 200, closed.text
    blocked = client.patch(
        url,
        headers=actor["headers"],
        json={
            "expected_current_version": 4,
            "renewal_due_date": "2028-05-31",
            "reason": "Cannot reactivate terminal source",
        },
    )
    assert blocked.status_code == 404
    with get_session_factory()() as session:
        record = session.get(IpDocketRecord, docket_id)
        assert record.status == "closed" and not record.is_active and record.current_version == 4


@pytest.mark.parametrize(
    ("name", "content", "expected"),
    [
        ("source.pdf", b"not supported", 415),
        ("source.xlsx", b"not a zip", 422),
        ("source.csv", b"", 422),
        ("source.csv", b"Trademark\n" + b"\n" * 1001, 422),
        ("source.csv", b"x" * (10 * 1024 * 1024 + 1), 413),
    ],
    ids=["unsupported-format", "bad-xlsx", "empty-file", "row-limit", "byte-limit"],
)
def test_file_failures_remain_typed_http_errors(client, name, content, expected):
    actor = _actor(client, configure=False)
    response = client.post(
        "/api/ip/imports/upload",
        headers=actor["headers"],
        files={"file": (name, content, "application/octet-stream")},
    )
    assert response.status_code == expected
    assert response.json()["status"] == expected


def test_unexpected_upload_exception_is_not_network_error_or_sensitive_log(
    client, monkeypatch, caplog
):
    actor = _actor(client)

    def crash(*args, **kwargs):
        raise RuntimeError("sensitive applicant must not be logged")

    monkeypatch.setattr("caseops_api.api.routes.ip_operations.create_ip_import_job", crash)
    response = client.post(
        "/api/ip/imports/upload",
        headers={
            **actor["headers"],
            "Origin": "http://localhost:3000",
            "X-Request-ID": "ip-upload-safe-500",
        },
        files={"file": ("source.csv", b"Trademark\nTest mark", "text/csv")},
    )
    assert response.status_code == 500
    assert response.json()["code"] == "ip_import_unexpected_failure"
    assert response.json()["request_id"] == "ip-upload-safe-500"
    assert response.headers["access-control-allow-origin"] == "http://localhost:3000"
    assert "sensitive applicant" not in response.text + caplog.text
    assert "RuntimeError" in caplog.text


def test_preview_duplicate_lookup_is_batched(client):
    actor = _actor(client)
    statements = []

    def observe(conn, cursor, statement, parameters, context, executemany):
        compiled = getattr(context, "compiled", None)
        statements.append(
            {
                "statement": statement,
                "read_only_select": is_read_only_select(getattr(compiled, "statement", None)),
            }
        )

    engine = get_engine()
    event.listen(engine, "before_cursor_execute", observe)
    try:
        _upload(
            client,
            actor,
            ["Trademark", "Application No."],
            [[f"Batch mark {i}", f"BATCH/{i}"] for i in range(100)],
        )
    finally:
        event.remove(engine, "before_cursor_execute", observe)
    for table in ("ip_identifiers", "ip_workspace_configurations"):
        matching = [
            row for row in statements if f"from {table}" in row["statement"].lower()
        ]
        assert len(matching) == 1, statements
        assert all(row["read_only_select"] for row in matching), statements


@pytest.mark.postgres
def test_postgres_late_duplicate_is_skipped(isolated_postgres_client):
    test_distinct_applications_with_same_mark_survive_and_late_duplicate_is_skipped(
        isolated_postgres_client,
    )


@pytest.mark.postgres
def test_postgres_concurrent_different_jobs_create_identifier_once(
    isolated_postgres_client, http_pg_engine, monkeypatch
):
    client = isolated_postgres_client
    actor = _actor(client)
    first = _upload(client, actor, ["Trademark", "Application No."], [["Concurrent", "CON/1"]])
    second = _upload(client, actor, ["Trademark", "Application No."], [["Concurrent", "CON 1"]])
    entered, release, contender = Event(), Event(), Event()
    real = ip_imports._materialize_import_row

    def paused(*args, **kwargs):
        entered.set()
        assert release.wait(5)
        return real(*args, **kwargs)

    monkeypatch.setattr(ip_imports, "_materialize_import_row", paused)

    def second_commit():
        contender.set()
        return _commit(client, actor, second, "concurrent-second-key")

    with ThreadPoolExecutor(max_workers=2) as pool:
        winner = pool.submit(_commit, client, actor, first, "concurrent-first-key")
        assert entered.wait(5)
        loser = pool.submit(second_commit)
        assert contender.wait(5)
        try:
            # Prove the contender reached this tenant's native advisory fence,
            # not merely that a future has not yet finished authentication.
            deadline = monotonic() + 3
            blocked = None
            resource = f"ip-import-tenant:{actor['company']['id']}"
            with http_pg_engine.connect() as observer:
                while monotonic() < deadline:
                    blocked = observer.execute(
                        text("""
                        SELECT waiting.pid, owner.pid
                        FROM pg_locks waiting JOIN pg_locks owner
                          ON owner.locktype = waiting.locktype
                         AND owner.classid = waiting.classid AND owner.objid = waiting.objid
                         AND owner.objsubid = waiting.objsubid AND owner.granted
                        WHERE waiting.locktype = 'advisory' AND NOT waiting.granted
                          AND waiting.classid =
                            ((hashtextextended(:resource, 0) >> 32) & 4294967295)
                          AND waiting.objid = (hashtextextended(:resource, 0) & 4294967295)
                          AND owner.pid = ANY(pg_blocking_pids(waiting.pid))
                    """),
                        {"resource": resource},
                    ).first()
                    if blocked is not None:
                        break
                    sleep(0.02)
            assert blocked is not None, "Missing directed native import blocker"
            assert blocked[0] != blocked[1]
        finally:
            release.set()
        results = [winner.result(10), loser.result(10)]
    assert sorted(result["job"]["committed_rows"] for result in results) == [0, 1]
    assert len(client.get("/api/ip/portfolio", headers=actor["headers"]).json()["rows"]) == 1
