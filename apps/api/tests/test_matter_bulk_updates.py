from __future__ import annotations

import csv
import io
import zipfile
from datetime import date

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import select

from caseops_api.db.models import AuditEvent, MatterHearing
from caseops_api.db.session import get_session_factory
from caseops_api.services import matter_access, matter_bulk_updates, matters
from caseops_api.services.matter_bulk_updates import HEADERS
from tests.test_auth_company import auth_headers, bootstrap_company

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _workbook_bytes(rows: list[list[object]], headers: list[str] = HEADERS) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.mark.parametrize("format_name", ["csv", "xlsx"])
def test_bulk_update_accepts_international_contact_number(format_name: str) -> None:
    values = [""] * len(HEADERS)
    values[HEADERS.index("Matter Code")] = "BULK-PHONE-1"
    values[HEADERS.index("Client Contact Number")] = "+919876543210"
    if format_name == "xlsx":
        content = _workbook_bytes([values])
    else:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerows([HEADERS, values])
        content = output.getvalue().encode("utf-8")
    _, rows, _ = matter_bulk_updates._parse_rows(content, f"updates.{format_name}")
    assert rows[0][1]["Client Contact Number"] == "+919876543210"


@pytest.mark.parametrize("unsafe_value", ["+SUM(1,2)", "=1+1", "@cmd"])
def test_bulk_update_still_rejects_formula_like_contact(unsafe_value: str) -> None:
    values = [""] * len(HEADERS)
    values[HEADERS.index("Matter Code")] = "BULK-PHONE-1"
    values[HEADERS.index("Client Contact Number")] = unsafe_value
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerows([HEADERS, values])
    with pytest.raises(HTTPException, match="Formula-like"):
        matter_bulk_updates._parse_rows(output.getvalue().encode("utf-8"), "updates.csv")


def test_bulk_update_rejects_xlsx_formula_node_in_contact_column() -> None:
    values = [""] * len(HEADERS)
    values[HEADERS.index("Matter Code")] = "BULK-PHONE-1"
    values[HEADERS.index("Client Contact Number")] = "=1+1"
    with pytest.raises(HTTPException, match="Spreadsheet formulas"):
        matter_bulk_updates._parse_rows(_workbook_bytes([values]), "updates.xlsx")


def test_bulk_update_previews_and_applies_existing_matter_only(client: TestClient) -> None:
    bootstrap_payload = bootstrap_company(client)
    token = str(bootstrap_payload["access_token"])
    matter_response = client.post(
        "/api/matters/",
        headers=auth_headers(token),
        json={
            "title": "Original bulk title",
            "matter_code": "BULK-UPDATE-1",
            "practice_area": "Commercial",
            "forum_level": "high_court",
            "status": "active",
        },
    )
    assert matter_response.status_code == 200, matter_response.text

    values = [""] * len(HEADERS)
    values[0] = "Updated bulk title"
    values[1] = "BULK-UPDATE-1"
    values[HEADERS.index("Matter Description")] = "Updated from the reviewed workbook"
    invalid = [""] * len(HEADERS)
    invalid[0] = "Must not be created"
    invalid[1] = "DOES-NOT-EXIST"
    content = _workbook_bytes([values, invalid])

    preview_response = client.post(
        "/api/matters/bulk-update/preview",
        headers=auth_headers(token),
        files={"file": ("updates.xlsx", content, XLSX_TYPE)},
    )
    assert preview_response.status_code == 200, preview_response.text
    preview = preview_response.json()
    assert preview["summary"] == {
        "total_rows": 2,
        "matched_rows": 1,
        "changed_rows": 1,
        "unchanged_rows": 0,
        "invalid_rows": 1,
    }

    apply_response = client.post(
        "/api/matters/bulk-update/apply",
        headers=auth_headers(token),
        data={"preview_token": preview["preview_token"]},
        files={"file": ("updates.xlsx", content, XLSX_TYPE)},
    )
    assert apply_response.status_code == 200, apply_response.text
    assert apply_response.json()["applied_rows"] == 1
    assert apply_response.json()["failed_rows"] == 0
    assert apply_response.json()["operation_id"]

    updated = client.get(
        f"/api/matters/{matter_response.json()['id']}",
        headers=auth_headers(token),
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["title"] == "Updated bulk title"
    assert updated.json()["description"] == "Updated from the reviewed workbook"
    history = client.get("/api/matters/bulk-update/history", headers=auth_headers(token))
    assert history.status_code == 200, history.text
    assert history.json()["total"] == 1
    assert history.json()["operations"][0]["applied_rows"] == 1
    assert history.json()["operations"][0]["invalid_rows"] == 1
    assert history.json()["operations"][0]["failed_rows"] == 0
    assert apply_response.json()["skipped_rows"] == 1


def test_bulk_update_cannot_change_lifecycle_status(client: TestClient) -> None:
    bootstrap_payload = bootstrap_company(client)
    token = str(bootstrap_payload["access_token"])
    matter_response = client.post(
        "/api/matters/",
        headers=auth_headers(token),
        json={
            "title": "Lifecycle bulk matter",
            "matter_code": "BULK-LIFECYCLE-1",
            "practice_area": "Civil",
            "forum_level": "high_court",
            "status": "active",
        },
    )
    assert matter_response.status_code == 200, matter_response.text
    legacy_headers = [*HEADERS[:4], "Matter Status", *HEADERS[4:]]
    values = [""] * len(legacy_headers)
    values[legacy_headers.index("Matter Code")] = "BULK-LIFECYCLE-1"
    values[legacy_headers.index("Matter Status")] = "Disposed"
    response = client.post(
        "/api/matters/bulk-update/preview",
        headers=auth_headers(token),
        files={"file": ("updates.xlsx", _workbook_bytes([values], legacy_headers), XLSX_TYPE)},
    )
    assert response.status_code == 400, response.text
    assert "Matter Status" not in HEADERS
    persisted = client.get(
        f"/api/matters/{matter_response.json()['id']}", headers=auth_headers(token)
    )
    assert persisted.json()["status"] == "active"


def test_bulk_update_hearing_date_uses_canonical_scheduling_service(client: TestClient) -> None:
    bootstrap_payload = bootstrap_company(client)
    token = str(bootstrap_payload["access_token"])
    matter_response = client.post(
        "/api/matters/",
        headers=auth_headers(token),
        json={
            "title": "Bulk hearing date",
            "matter_code": "BULK-HEARING-1",
            "practice_area": "Civil",
            "forum_level": "high_court",
            "status": "active",
        },
    )
    assert matter_response.status_code == 200, matter_response.text
    values = [""] * len(HEADERS)
    values[HEADERS.index("Matter Code")] = "BULK-HEARING-1"
    values[HEADERS.index("Next Hearing Date")] = "2026-10-15"
    content = _workbook_bytes([values])
    preview = client.post(
        "/api/matters/bulk-update/preview",
        headers=auth_headers(token),
        files={"file": ("updates.xlsx", content, XLSX_TYPE)},
    )
    assert preview.status_code == 200, preview.text
    date_change = preview.json()["rows"][0]["changes"]["next_hearing_on"]
    assert date_change == {"old": None, "new": "2026-10-15"}
    with get_session_factory()() as session:
        assert session.scalar(
            select(MatterHearing.id).where(MatterHearing.matter_id == matter_response.json()["id"])
        ) is None
    applied = client.post(
        "/api/matters/bulk-update/apply",
        headers=auth_headers(token),
        data={"preview_token": preview.json()["preview_token"]},
        files={"file": ("updates.xlsx", content, XLSX_TYPE)},
    )
    assert applied.status_code == 200, applied.text
    assert applied.json()["applied_rows"] == 1
    with get_session_factory()() as session:
        hearings = list(
            session.scalars(
                select(MatterHearing).where(
                    MatterHearing.matter_id == matter_response.json()["id"],
                    MatterHearing.hearing_on == date.fromisoformat("2026-10-15"),
                )
            )
        )
    assert len(hearings) == 1
    assert hearings[0].status == "scheduled"


def test_csv_template_duplicate_identity_rejection_and_history(client: TestClient) -> None:
    bootstrap_payload = bootstrap_company(client)
    token = str(bootstrap_payload["access_token"])
    matter = client.post(
        "/api/matters/",
        headers=auth_headers(token),
        json={
            "title": "CSV baseline",
            "matter_code": "BULK-CSV-1",
            "practice_area": "Civil",
            "forum_level": "high_court",
            "status": "active",
        },
    )
    assert matter.status_code == 200, matter.text
    template = client.get(
        "/api/matters/bulk-update/template?format=csv", headers=auth_headers(token)
    )
    assert template.status_code == 200, template.text
    assert next(csv.reader(io.StringIO(template.text))) == HEADERS

    values = [""] * len(HEADERS)
    values[HEADERS.index("Matter Code")] = "BULK-CSV-1"
    values[HEADERS.index("Matter Title")] = "CSV updated title"
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\r\n")
    writer.writerow(HEADERS)
    writer.writerow(values)
    content = output.getvalue().encode()
    preview = client.post(
        "/api/matters/bulk-update/preview",
        headers=auth_headers(token),
        files={"file": ("updates.csv", content, "text/csv")},
    )
    assert preview.status_code == 200, preview.text
    apply = client.post(
        "/api/matters/bulk-update/apply",
        headers=auth_headers(token),
        data={"preview_token": preview.json()["preview_token"]},
        files={"file": ("updates.csv", content, "text/csv")},
    )
    assert apply.status_code == 200, apply.text
    assert apply.json()["total_rows"] == 1
    assert apply.json()["applied_rows"] == 1
    assert apply.json()["skipped_rows"] == 0
    updated = client.get(f"/api/matters/{matter.json()['id']}", headers=auth_headers(token))
    assert updated.json()["title"] == "CSV updated title"
    history = client.get("/api/matters/bulk-update/history", headers=auth_headers(token))
    assert history.json()["operations"][0]["format"] == "csv"


def test_duplicate_matter_codes_are_both_rejected_before_apply(client: TestClient) -> None:
    bootstrap_payload = bootstrap_company(client)
    token = str(bootstrap_payload["access_token"])
    matter = client.post(
        "/api/matters/",
        headers=auth_headers(token),
        json={
            "title": "Duplicate baseline",
            "matter_code": "BULK-DUP-1",
            "practice_area": "Civil",
            "forum_level": "high_court",
            "status": "active",
        },
    )
    assert matter.status_code == 200, matter.text
    rows = []
    for title in ("First title", "Second title"):
        row = [""] * len(HEADERS)
        row[HEADERS.index("Matter Code")] = "BULK-DUP-1"
        row[HEADERS.index("Matter Title")] = title
        rows.append(row)
    content = _workbook_bytes(rows)
    preview = client.post(
        "/api/matters/bulk-update/preview",
        headers=auth_headers(token),
        files={"file": ("duplicates.xlsx", content, XLSX_TYPE)},
    )
    assert preview.status_code == 200, preview.text
    assert [row["status"] for row in preview.json()["rows"]] == ["invalid", "invalid"]
    assert all("appears more than once" in row["errors"][0] for row in preview.json()["rows"])
    updated = client.get(f"/api/matters/{matter.json()['id']}", headers=auth_headers(token))
    assert updated.json()["title"] == "Duplicate baseline"


def test_bulk_preview_conceals_ethical_wall_matter_like_a_missing_code(
    client: TestClient,
) -> None:
    boot = bootstrap_company(client)
    token = str(boot["access_token"])
    hidden = client.post(
        "/api/matters/",
        headers=auth_headers(token),
        json={
            "title": "Confidential matter",
            "matter_code": "BULK-HIDDEN-1",
            "practice_area": "Civil",
            "forum_level": "high_court",
            "status": "active",
        },
    )
    assert hidden.status_code == 200, hidden.text
    invited = client.post(
        "/api/companies/current/users",
        headers=auth_headers(token),
        json={
            "full_name": "Bulk Member",
            "email": "bulk-member@example.com",
            "role": "member",
            "password": "BulkMemberPass123!",
        },
    )
    assert invited.status_code == 200, invited.text
    restricted = client.post(
        f"/api/matters/{hidden.json()['id']}/access/restricted",
        headers=auth_headers(token),
        json={"restricted": True},
    )
    assert restricted.status_code == 200, restricted.text
    login = client.post(
        "/api/auth/login",
        json={
            "company_slug": "aster-legal",
            "email": "bulk-member@example.com",
            "password": "BulkMemberPass123!",
        },
    )
    assert login.status_code == 200, login.text
    rows = []
    for code in ("BULK-HIDDEN-1", "BULK-DOES-NOT-EXIST"):
        values = [""] * len(HEADERS)
        values[HEADERS.index("Matter Code")] = code
        values[HEADERS.index("Matter Title")] = "Attempted change"
        rows.append(values)
    preview = client.post(
        "/api/matters/bulk-update/preview",
        headers=auth_headers(str(login.json()["access_token"])),
        files={"file": ("updates.xlsx", _workbook_bytes(rows), XLSX_TYPE)},
    )
    assert preview.status_code == 200, preview.text
    concealed, missing = preview.json()["rows"]
    for row in (concealed, missing):
        assert row["status"] == "invalid"
        assert row["matter_id"] is None
        assert row["expected_updated_at"] is None
        assert row["changes"] == {}
    assert concealed["errors"] == missing["errors"]
    with get_session_factory()() as session:
        denied = session.scalar(
            select(AuditEvent.id).where(
                AuditEvent.company_id == boot["company"]["id"],
                AuditEvent.target_id == hidden.json()["id"],
                AuditEvent.action == "access_denied",
            )
        )
    assert denied is not None


def test_bulk_preview_shows_canonical_court_lineage_removals(client: TestClient) -> None:
    token = str(bootstrap_company(client)["access_token"])
    created = client.post(
        "/api/matters/",
        headers=auth_headers(token),
        json={
            "title": "Catalog court baseline",
            "matter_code": "BULK-COURT-1",
            "practice_area": "Civil",
            "forum_level": "lower_court",
            "court_name": "saket-court",
            "forum_district": "South Delhi",
            "status": "active",
        },
    )
    assert created.status_code == 200, created.text
    assert created.json()["forum_catalog_entry_id"]
    values = [""] * len(HEADERS)
    values[HEADERS.index("Matter Code")] = "BULK-COURT-1"
    values[HEADERS.index("Court")] = "Imaginary Example Court"
    content = _workbook_bytes([values])
    preview = client.post(
        "/api/matters/bulk-update/preview",
        headers=auth_headers(token),
        files={"file": ("updates.xlsx", content, XLSX_TYPE)},
    )
    assert preview.status_code == 200, preview.text
    changes = preview.json()["rows"][0]["changes"]
    assert changes["forum_catalog_entry_id"] == {
        "old": created.json()["forum_catalog_entry_id"],
        "new": None,
    }
    assert changes["court_name"]["new"] == "Imaginary Example Court"
    assert changes["forum_state"]["new"] is None
    applied = client.post(
        "/api/matters/bulk-update/apply",
        headers=auth_headers(token),
        data={"preview_token": preview.json()["preview_token"]},
        files={"file": ("updates.xlsx", content, XLSX_TYPE)},
    )
    assert applied.status_code == 200, applied.text
    persisted = client.get(
        f"/api/matters/{created.json()['id']}", headers=auth_headers(token)
    ).json()
    for field_name, change in changes.items():
        assert persisted[field_name] == change["new"]


def test_bulk_preview_rejects_out_of_team_owner_before_apply(client: TestClient) -> None:
    boot = bootstrap_company(client)
    token = str(boot["access_token"])
    team = client.post(
        "/api/teams/", headers=auth_headers(token), json={"name": "Litigation", "slug": "lit"}
    )
    assert team.status_code == 201, team.text
    other_team = client.post(
        "/api/teams/", headers=auth_headers(token), json={"name": "IP", "slug": "ip"}
    )
    assert other_team.status_code == 201, other_team.text
    outsider = client.post(
        "/api/companies/current/users",
        headers=auth_headers(token),
        json={
            "full_name": "Outside Team",
            "email": "outside-team@example.com",
            "role": "member",
            "password": "OutsideTeamPass123!",
        },
    )
    assert outsider.status_code == 200, outsider.text
    added = client.post(
        f"/api/teams/{other_team.json()['id']}/members",
        headers=auth_headers(token),
        json={"membership_id": outsider.json()["membership_id"]},
    )
    assert added.status_code == 200, added.text
    matter = client.post(
        "/api/matters/",
        headers=auth_headers(token),
        json={
            "title": "Team-scoped bulk target",
            "matter_code": "BULK-TEAM-1",
            "practice_area": "Civil",
            "forum_level": "high_court",
            "status": "active",
        },
    )
    assert matter.status_code == 200, matter.text
    assigned = client.patch(
        f"/api/matters/{matter.json()['id']}",
        headers=auth_headers(token),
        json={"team_id": team.json()["id"], "expected_updated_at": matter.json()["updated_at"]},
    )
    assert assigned.status_code == 200, assigned.text
    enabled = client.put("/api/teams/scoping", headers=auth_headers(token), json={"enabled": True})
    assert enabled.status_code == 200, enabled.text
    values = [""] * len(HEADERS)
    values[HEADERS.index("Matter Code")] = "BULK-TEAM-1"
    values[HEADERS.index("Matter Owner")] = "outside-team@example.com"
    preview = client.post(
        "/api/matters/bulk-update/preview",
        headers=auth_headers(token),
        files={"file": ("updates.xlsx", _workbook_bytes([values]), XLSX_TYPE)},
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["rows"][0]["status"] == "invalid"
    assert "Matter owner must belong to the assigned team." in preview.json()["rows"][0]["errors"]


def test_bulk_rejects_wide_xlsx_and_csv_rows(client: TestClient) -> None:
    token = str(bootstrap_company(client)["access_token"])
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    sheet.cell(row=2, column=16384, value="unreviewed")
    output = io.BytesIO()
    workbook.save(output)
    wide = client.post(
        "/api/matters/bulk-update/preview",
        headers=auth_headers(token),
        files={"file": ("wide.xlsx", output.getvalue(), XLSX_TYPE)},
    )
    assert wide.status_code == 400, wide.text
    assert "columns outside" in wide.text
    csv_output = io.StringIO(newline="")
    writer = csv.writer(csv_output)
    writer.writerow(HEADERS)
    writer.writerow(["ignored"] * (len(HEADERS) + 1))
    csv_response = client.post(
        "/api/matters/bulk-update/preview",
        headers=auth_headers(token),
        files={"file": ("wide.csv", csv_output.getvalue().encode(), "text/csv")},
    )
    assert csv_response.status_code == 400, csv_response.text


def test_bulk_preview_rejects_disposed_matter_without_writing(client: TestClient) -> None:
    token = str(bootstrap_company(client)["access_token"])
    created = client.post(
        "/api/matters/",
        headers=auth_headers(token),
        json={
            "title": "Terminal bulk target",
            "matter_code": "BULK-TERMINAL-1",
            "practice_area": "Civil",
            "forum_level": "high_court",
            "status": "active",
        },
    )
    assert created.status_code == 200, created.text
    disposed = client.patch(
        f"/api/matters/{created.json()['id']}/lifecycle/status",
        headers=auth_headers(token),
        json={
            "to_status": "disposed",
            "expected_from_status": "active",
            "expected_updated_at": created.json()["updated_at"],
            "reason": "Final judgment entered and engagement completed",
        },
    )
    assert disposed.status_code == 200, disposed.text
    values = [""] * len(HEADERS)
    values[HEADERS.index("Matter Code")] = "BULK-TERMINAL-1"
    values[HEADERS.index("Matter Title")] = "Improper reopening"
    preview = client.post(
        "/api/matters/bulk-update/preview",
        headers=auth_headers(token),
        files={"file": ("updates.xlsx", _workbook_bytes([values]), XLSX_TYPE)},
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["rows"][0]["status"] == "invalid"
    assert "immutable" in str(preview.json()["rows"][0]["errors"])
    persisted = client.get(f"/api/matters/{created.json()['id']}", headers=auth_headers(token))
    assert persisted.json()["status"] == "disposed"
    assert persisted.json()["title"] == "Terminal bulk target"


def test_bulk_rejects_compressed_xlsx_expansion_before_parsing(client: TestClient) -> None:
    token = str(bootstrap_company(client)["access_token"])
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "xl/sharedStrings.xml", b"A" * (matter_bulk_updates.MAX_XLSX_PART_BYTES + 1)
        )
    response = client.post(
        "/api/matters/bulk-update/preview",
        headers=auth_headers(token),
        files={"file": ("compressed.xlsx", output.getvalue(), XLSX_TYPE)},
    )
    assert response.status_code == 400, response.text
    assert "oversized part" in response.text


def test_bulk_apply_rejects_late_court_resolution_drift(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    token = str(bootstrap_company(client)["access_token"])
    created = client.post(
        "/api/matters/",
        headers=auth_headers(token),
        json={
            "title": "Stable title",
            "matter_code": "BULK-DRIFT-1",
            "practice_area": "Civil",
            "forum_level": "high_court",
            "status": "active",
        },
    )
    assert created.status_code == 200, created.text
    values = [""] * len(HEADERS)
    values[HEADERS.index("Matter Code")] = "BULK-DRIFT-1"
    values[HEADERS.index("Matter Title")] = "Changed title"
    values[HEADERS.index("Court")] = "First reviewed court"
    content = _workbook_bytes([values])
    preview = client.post(
        "/api/matters/bulk-update/preview",
        headers=auth_headers(token),
        files={"file": ("updates.xlsx", content, XLSX_TYPE)},
    )
    assert preview.status_code == 200, preview.text
    original = matters._resolve_forum_selection
    calls = 0

    def changed_resolution(*args, **kwargs):
        nonlocal calls
        calls += 1
        selection = original(*args, **kwargs)
        if calls == 2:
            return {**selection, "court_name": "Unreviewed late court"}
        return selection

    monkeypatch.setattr(matters, "_resolve_forum_selection", changed_resolution)
    applied = client.post(
        "/api/matters/bulk-update/apply",
        headers=auth_headers(token),
        data={"preview_token": preview.json()["preview_token"]},
        files={"file": ("updates.xlsx", content, XLSX_TYPE)},
    )
    assert applied.status_code == 409, applied.text
    assert calls == 2
    persisted = client.get(f"/api/matters/{created.json()['id']}", headers=auth_headers(token))
    assert persisted.json()["title"] == "Stable title"
    assert persisted.json()["court_name"] is None


def _assert_bulk_access_change_rolls_back_batch(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    boot = bootstrap_company(client)
    token = str(boot["access_token"])
    matters = []
    rows = []
    for number in (1, 2):
        created = client.post(
            "/api/matters/",
            headers=auth_headers(token),
            json={
                "title": f"Before {number}",
                "matter_code": f"BULK-ATOMIC-{number}",
                "practice_area": "Civil",
                "forum_level": "high_court",
                "status": "active",
            },
        )
        assert created.status_code == 200, created.text
        matters.append(created.json())
        row = [""] * len(HEADERS)
        row[HEADERS.index("Matter Code")] = f"BULK-ATOMIC-{number}"
        row[HEADERS.index("Matter Title")] = f"After {number}"
        rows.append(row)
    content = _workbook_bytes(rows)
    preview = client.post(
        "/api/matters/bulk-update/preview",
        headers=auth_headers(token),
        files={"file": ("updates.xlsx", content, XLSX_TYPE)},
    )
    assert preview.status_code == 200, preview.text
    blocked_id = max(matter["id"] for matter in matters)
    state = {"first_applied": False}
    real_can_access = matter_access.can_access
    real_update = matter_bulk_updates.update_matter

    def changed_access(session, *, context, matter):
        if state["first_applied"] and matter.id == blocked_id:
            return False
        return real_can_access(session, context=context, matter=matter)

    def update_then_revoke(*args, **kwargs):
        result = real_update(*args, **kwargs)
        if not args[0].in_nested_transaction():
            state["first_applied"] = True
        return result

    monkeypatch.setattr(matter_access, "can_access", changed_access)
    monkeypatch.setattr(matter_bulk_updates, "can_access", changed_access)
    monkeypatch.setattr(matter_bulk_updates, "update_matter", update_then_revoke)
    applied = client.post(
        "/api/matters/bulk-update/apply",
        headers=auth_headers(token),
        data={"preview_token": preview.json()["preview_token"]},
        files={"file": ("updates.xlsx", content, XLSX_TYPE)},
    )
    assert applied.status_code == 409, applied.text
    assert state["first_applied"]
    assert "No rows were updated" in applied.text
    monkeypatch.undo()
    for matter in matters:
        persisted = client.get(f"/api/matters/{matter['id']}", headers=auth_headers(token))
        assert persisted.status_code == 200, persisted.text
        assert persisted.json()["title"] == matter["title"]
    history = client.get("/api/matters/bulk-update/history", headers=auth_headers(token))
    assert history.status_code == 200, history.text
    assert history.json()["operations"][0]["status"] == "stale"
    assert history.json()["operations"][0]["applied_rows"] == 0
    with get_session_factory()() as session:
        denied = session.scalar(
            select(AuditEvent.id).where(
                AuditEvent.company_id == boot["company"]["id"],
                AuditEvent.target_id == blocked_id,
                AuditEvent.action == "access_denied",
            )
        )
    assert denied is not None


def test_bulk_access_change_rolls_back_batch(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    _assert_bulk_access_change_rolls_back_batch(client, monkeypatch)


@pytest.mark.postgres
def test_bulk_access_change_rolls_back_batch_on_postgres(
    isolated_postgres_client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    _assert_bulk_access_change_rolls_back_batch(isolated_postgres_client, monkeypatch)
